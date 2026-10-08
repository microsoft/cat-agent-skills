#!/usr/bin/env python3
"""Build the step-by-step Word document from steps.json (output of linearize.py).

Usage:  python build_docx.py steps.json -o process.docx [--title "Customer onboarding"] [--source diagram.svg]

Requires python-docx. Verifies that every numbered step is written exactly once and exits 2 if not.
"""
import argparse
import json
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CALLOUT_FILL = 'E8EEF7'


def shade(paragraph, fill):
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    ppr.append(shd)


def callout(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    shade(p, CALLOUT_FILL)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(8)
    return p


def field(doc, label, value):
    p = doc.add_paragraph()
    p.add_run(f'{label}: ').bold = True
    p.add_run(value)
    p.paragraph_format.space_after = Pt(2)


class Writer:
    def __init__(self, doc, data):
        self.doc, self.data = doc, data
        self.steps = {s['number']: s for s in data['steps']}
        self.parallel = data.get('parallel', {})
        self.rendered = {}

    def branch_lines(self, step):
        n = step['number']
        if step['kind'] == 'end':
            return ['End of process.']
        if str(n) in self.parallel:
            pw = self.parallel[str(n)]['pathways']
            return ['Forks into %d pathways that run concurrently: %s.' % (
                len(pw), ', '.join(f'Pathway {p["letter"]} (Step {p["root"]})' for p in pw))]
        br = step['branches']
        if not br:
            return ['No outgoing flow: the path ends here.']
        if len(br) == 1:
            b = br[0]
            if b['kind'] == 'link':
                return [f'Hands off via link "{b["link_name"]}" to Step {b["target"]}.']
            if b['loopback']:
                return [f'Loops back to Step {b["target"]}.']
            return [f'Continues to Step {b["target"]}.']
        lines = []
        for b in br:
            label = b['label'] or '(no label)'
            tail = f'loops back to Step {b["target"]}' if b['loopback'] else f'Step {b["target"]}'
            lines.append(f'If "{label}": {tail}')
        return lines

    def step(self, n):
        s = self.steps[n]
        self.rendered[n] = self.rendered.get(n, 0) + 1
        self.doc.add_heading(f'Step {n} \u2014 {s["type_label"]}: {s["name"]}', level=2 + 2 * self._depth)
        if s['kind'] == 'gateway':
            field(self.doc, 'Gateway type', s['gateway_type'])
        if s.get('performer') and s['kind'] == 'activity':
            field(self.doc, 'Performer', s['performer'])
        if s.get('description'):
            field(self.doc, 'Description', s['description'])
        if s.get('annotation'):
            field(self.doc, 'Note', s['annotation'])
        for line in self.branch_lines(s):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_after = Pt(2)
            p.add_run('\u2192 ').bold = True
            p.add_run(line)

    _depth = 0

    def sequence(self, numbers):
        for n in sorted(numbers):
            if n in self.rendered:
                continue
            self.step(n)
            fork = self.parallel.get(str(n))
            if not fork:
                continue
            callout(self.doc, '\u25C6 PARALLEL GATEWAY (+ symbol) \u2014 forks into %d pathways that run concurrently.'
                    % len(fork['pathways']))
            for p in fork['pathways']:
                h = self.doc.add_heading(f'Pathway {p["letter"]}: {p["name"]}', level=3 + 2 * self._depth)
                for r in h.runs:
                    r.bold = True
                self._depth += 1
                self.sequence(p['steps'])
                self._depth -= 1
            if fork['resume']:
                callout(self.doc, f'\u25C6 PARALLEL GATEWAY JOIN \u2014 all pathways complete. '
                                  f'Process resumes at Step {fork["resume"]}.')
            else:
                callout(self.doc, '\u25C6 PARALLEL GATEWAY JOIN \u2014 the pathways do not rejoin; '
                                  'each ends independently.')

    def verify(self):
        missing = sorted(set(self.steps) - set(self.rendered))
        dupes = sorted(n for n, c in self.rendered.items() if c > 1)
        return missing, dupes


def build(data, out, title, source):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(10.5)
    doc.core_properties.title = f'Process documentation: {title}'
    for name, size in (('Heading 3', 13), ('Heading 4', 11.5), ('Heading 5', 11.5), ('Heading 6', 10.5)):
        st = doc.styles[name]
        st.font.italic, st.font.size = False, Pt(size)

    doc.add_heading(f'Process documentation: {title}', level=0)
    if source:
        p = doc.add_paragraph()
        p.add_run(f'Source diagram: {source}').italic = True
    st = data['stats']
    doc.add_heading('Overview', level=1)
    decisions = sum(1 for s in data['steps'] if s['kind'] == 'gateway' and not s['gateway_type'].startswith('Parallel'))
    for line in (f'{st["steps"]} numbered steps (tasks, events and branching gateways).',
                 f'{decisions} decision point(s); {st["parallel_forks"]} parallel fork(s); {st["loops"]} loop(s).',
                 'Step numbers follow the diagram left to right and are global: parallel pathways keep them.'):
        doc.add_paragraph(line, style='List Bullet')

    doc.add_heading('Step-by-step description', level=1)
    w = Writer(doc, data)
    w.sequence(list(w.steps))
    missing, dupes = w.verify()

    if data.get('warnings'):
        doc.add_heading('Items to confirm with the process owner', level=1)
        for msg in dict.fromkeys(data['warnings']):
            doc.add_paragraph(msg, style='List Bullet')
    doc.save(out)
    return missing, dupes


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('steps')
    ap.add_argument('-o', '--output', default='process.docx')
    ap.add_argument('--title', default='BPMN process')
    ap.add_argument('--source', default='')
    args = ap.parse_args()
    with open(args.steps, encoding='utf-8') as fh:
        data = json.load(fh)
    missing, dupes = build(data, args.output, args.title, args.source)
    print(f'wrote {args.output}: {len(data["steps"])} steps', file=sys.stderr)
    if missing or dupes:
        print(f'FATAL: steps missing from document: {missing}; written more than once: {dupes}', file=sys.stderr)
        return 2
    return 0


if __name__ == '__main__':
    sys.exit(main())
