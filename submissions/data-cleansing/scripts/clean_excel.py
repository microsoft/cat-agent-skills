#!/usr/bin/env python3
"""Clean Excel workbooks: fix what is safe to fix, quarantine the rest.

Requires pandas + openpyxl (preinstalled in the Copilot Studio code interpreter).
For each input workbook a new `<name>_cleaned.xlsx` is written. The original is never modified.

Output workbook sheets:
  Summary              per-sheet counts + warnings that need a human decision
  <sheet>              the cleaned rows
  <sheet> Rejected     rows removed from the cleaned data, original values + reason + source row
  Change Log           every automatic fix, with counts and an example
"""
from __future__ import annotations

import argparse
import math
import re
import sys
import warnings
from collections import Counter
from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd

warnings.filterwarnings("ignore")

EXCEL_EXTS = {".xlsx", ".xlsm", ".xls"}

CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
INVISIBLE_RE = re.compile(r"[\u200b\u200c\u200d\u2060\ufeff]")
WS_RE = re.compile(r"\s+")
ERROR_TOKENS = {"#n/a", "#div/0!", "#value!", "#ref!", "#name?", "#null!", "#num!"}
NULL_TOKENS = {"", "n/a", "nan", "null", "nil", "-", "--", "(blank)"} | ERROR_TOKENS

NUM_RE = re.compile(r"^[-+]?(\d+\.?\d*|\.\d+)([eE][-+]?\d+)?$")
THOUSANDS_RE = re.compile(r"^[-+]?\d{1,3}(,\d{3})+(\.\d+)?$")
LONG_INT_RE = re.compile(r"^[-+]?\d{16,}$")
ID_NAME_RE = re.compile(
    r"(^|[^a-z])(id|zip|postal|postcode|phone|tel|fax|sku|code|ssn|sin|account|acct|number|no|num)([^a-z]|$)",
    re.I,
)
# Heading hints are matched as whole words, never substrings ("Candidate ID" is not a date, "Account" is not a count).
DATE_WORDS = {"date", "datetime", "time", "timestamp", "dob", "birth", "birthday", "created", "updated", "modified"}
# headings that contain an identifier word but hold quantities/amounts (e.g. "Number of Units", "Account Balance")
QTY_WORDS = {"qty", "quantity", "units", "amount", "total", "price", "rate", "balance", "revenue", "cost", "count"}
QTY_PHRASE_RE = re.compile(r"\b(number|num|no)\.? of\b|#\s*of\b", re.I)
# "Email Status", "Email Opt In" ... describe an address, they do not contain one
NON_ADDRESS_WORDS = {"status", "opt", "optin", "optout", "verified", "valid", "validated", "flag", "sent", "bounced",
                     "bounce", "consent", "subscribed", "subscription", "preference", "pref", "domain", "type",
                     "source", "count", "date", "time", "reason", "enabled", "allowed", "confirmed", "marketing"}
# optional sign and ONE currency symbol (leading or trailing); digits/commas/dot must be contiguous (no inner spaces)
NUM_PARTS_RE = re.compile(
    r"^(?P<s1>[-+]?)\s*(?P<cur1>[$\u20ac\u00a3\u00a5]?)\s*(?P<s2>[-+]?)(?P<num>[\d,.]+(?:[eE][-+]?\d+)?)\s*(?P<cur2>[$\u20ac\u00a3\u00a5]?)$"
)
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
AMBIGUOUS_DATE_RE = re.compile(r"^\s*(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})\s*$")


# ---------------------------------------------------------------- helpers
def isnull(v) -> bool:
    if v is None or v is pd.NaT:
        return True
    return isinstance(v, float) and math.isnan(v)


def clean_cell(v):
    """Normalise one cell: strip junk characters/whitespace, turn placeholders into empty."""
    if isnull(v):
        return None
    if not isinstance(v, str):
        return v
    t = CONTROL_RE.sub("", INVISIBLE_RE.sub("", v))
    t = WS_RE.sub(" ", t).strip()
    return None if t.casefold() in NULL_TOKENS else t


def parse_number(v):
    """Return (value, is_percent) or None when v is not unambiguously numeric."""
    if isinstance(v, (bool, np.bool_)):
        return None
    if isinstance(v, (int, float, np.integer, np.floating)):
        x = float(v)
        return (x, False) if math.isfinite(x) else None
    if not isinstance(v, str):
        return None
    t = v.strip()
    neg = t.startswith("(") and t.endswith(")")
    if neg:
        t = t[1:-1].strip()
    pct = t.endswith("%")
    if pct:
        t = t[:-1].strip()
    m = NUM_PARTS_RE.match(t)
    if not m or (m["s1"] and m["s2"]) or (m["cur1"] and m["cur2"]):
        return None  # malformed, e.g. "1 2", "$1$2", "--5"
    sign = m["s1"] or m["s2"]
    if neg and sign:
        return None  # '(-200)' contradicts itself; flipping or ignoring the sign could reverse a balance, so leave it for review
    num = m["num"]
    if "," in num:
        if not THOUSANDS_RE.match(num):
            return None  # e.g. "1,5" (decimal comma) - too risky to guess
        num = num.replace(",", "")
    if not NUM_RE.match(num):
        return None
    x = float(sign + num)
    if not math.isfinite(x):
        return None  # "1e999" overflows to infinity, which would make distinct values compare equal
    if neg:
        x = -x
    return (x / 100.0 if pct else x), pct


def name_tokens(col: str) -> list:
    """Lower-case whole words of a heading, splitting camelCase and punctuation ('AccountCreatedDate' -> 3 words)."""
    return re.findall(r"[a-z]+", re.sub(r"([a-z])([A-Z])", r"\1 \2", col).lower())


def is_date_name(col: str) -> bool:
    return bool(DATE_WORDS & set(name_tokens(col)))


def is_qty_name(col: str) -> bool:
    return bool(QTY_WORDS & set(name_tokens(col))) or bool(QTY_PHRASE_RE.search(col))


def is_email_name(col: str) -> bool:
    tokens = set(name_tokens(col))
    return bool({"email", "emails", "mail"} & tokens) and not (NON_ADDRESS_WORDS & tokens)


class HeaderReviewNeeded(Exception):
    """The header row cannot be chosen safely; a person must confirm it."""


def header_row_for(sheet: str, specs):
    """--header-row accepts N (every sheet) and/or SHEET=N (one sheet); the sheet-specific value wins."""
    default = None
    for spec in specs or []:
        if "=" in spec:
            s, n = spec.rsplit("=", 1)
            if s.strip().casefold() == sheet.casefold():
                return int(n)
        else:
            default = int(spec)
    return default


TZ_RE = re.compile(r"(Z|[+-]\d{2}:?\d{2})\s*$")


def parse_dates(strings: pd.Series, dayfirst: bool) -> pd.Series:
    """Parse to timezone-naive timestamps. Timezone-aware values are converted to UTC, then made naive."""
    try:
        out = pd.to_datetime(strings, errors="coerce", format="mixed", dayfirst=dayfirst, utc=True)
    except (TypeError, ValueError):
        out = pd.to_datetime(strings, errors="coerce", dayfirst=dayfirst, utc=True)
    return out.dt.tz_localize(None)


def strip_illegal(v):
    return CONTROL_RE.sub("", v) if isinstance(v, str) else v


def short(v, n=60) -> str:
    s = repr(v) if isinstance(v, str) else str(v)
    return s if len(s) <= n else s[: n - 3] + "..."


def free_name(base: str, existing) -> str:
    """base, or 'base (2)', 'base (3)'... - whichever is not already a column name."""
    taken, cand, i = set(existing), base, 2
    while cand in taken:
        cand, i = f"{base} ({i})", i + 1
    return cand


def above_rows_frame(above_raw: pd.DataFrame, reason_col: str, row_col: str) -> pd.DataFrame:
    """Rows found above the header, laid out like a Rejected sheet (source row + reason + original values)."""
    out = above_raw.apply(lambda c: c.map(strip_illegal))
    out.insert(0, reason_col, "Row above the header (title/note) - not part of the data table")
    out.insert(0, row_col, out.index)
    return out


def unique_sheet_name(name: str, used: set) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", name)[:31] or "Sheet"
    cand, i = base, 2
    while cand.casefold() in used:
        suffix = f" ({i})"
        cand = base[: 31 - len(suffix)] + suffix
        i += 1
    used.add(cand.casefold())
    return cand


# ---------------------------------------------------------------- reading
def read_workbook(path: Path) -> dict:
    """Return {sheet name: DataFrame(object dtype)} indexed by the real Excel row number."""
    sheets = {}
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=True)  # cached formula results, not formulas
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if rows:
                sheets[ws.title] = pd.DataFrame(rows, index=range(1, len(rows) + 1), dtype=object)
    else:
        # keep_default_na=False: a literal "NA" (e.g. a region code) must reach clean_cell, which owns the placeholder policy
        for name, df in pd.read_excel(path, sheet_name=None, header=None, dtype=object, keep_default_na=False).items():
            df = df.astype(object)
            df.index = range(1, len(df) + 1)
            sheets[name] = df
    return sheets


# ---------------------------------------------------------------- cleaning
def clean_sheet(name: str, raw: pd.DataFrame, opts) -> dict | None:
    log, warns = [], []

    def note(column, action, count, example=""):
        if count:
            log.append({"Sheet": name, "Column": column, "Action": action, "Count": int(count), "Example": example})

    cleaned = raw.apply(lambda col: col.map(clean_cell))

    # --- locate the header row (skips title rows / blank lines above the table)
    nn = cleaned.notna().sum(axis=1)
    if nn.max() == 0:
        return None
    window = nn[nn > 0].iloc[:20]
    thr = max(1, math.ceil(0.6 * window.max()))
    header_idx = window.index[int((window.values >= thr).argmax())]
    explicit = header_row_for(name, opts.header_row)
    confirmed = False
    if explicit:
        if explicit in nn.index and nn.loc[explicit] > 0:
            header_idx, confirmed = explicit, True
        else:
            warns.append(f"[{name}] header row {explicit} is empty or outside the sheet - header was auto-detected instead.")
    above_idx = list(nn.index[(nn.index < header_idx) & (nn > 0)])
    multi_above = [i for i in above_idx if nn.loc[i] >= 2]
    if multi_above and not confirmed:
        # a multi-cell row above the detected header could itself be a header with empty cells (or the detected
        # "header" could be a data row), so guessing could mislabel every column - stop and ask instead
        raise HeaderReviewNeeded(
            f"[{name}] cannot tell which row is the header. Row {header_idx} looks like the header, but row(s) "
            f"{', '.join(map(str, multi_above))} above it also hold several values. Nothing was changed for this sheet. "
            f'Confirm the header row N and re-run with --header-row "{name}=N".'
        )
    note("(sheet)", "Moved title/note rows above the header to the Rejected sheet", len(above_idx))

    names, auto = [], []
    for j, v in enumerate(cleaned.loc[header_idx]):
        if isnull(v):  # pandas may hand back NaN rather than None for an empty header cell
            names.append(f"Column_{j + 1}")
            auto.append(True)
        else:
            names.append(str(int(v)) if isinstance(v, float) and v.is_integer() else str(v))
            auto.append(False)
    seen, taken = Counter(), set(names)  # reserve every existing header so suffixes never collide
    for j, n in enumerate(list(names)):
        seen[n] += 1
        if seen[n] > 1:
            k = seen[n]
            while f"{n}_{k}" in taken:
                k += 1
            names[j] = f"{n}_{k}"
            taken.add(names[j])
            note(n, "Renamed duplicate column header", 1, f"{n} -> {names[j]}")

    body = cleaned.loc[header_idx + 1:].copy()
    raw_body = raw.loc[header_idx + 1:].copy()
    body.columns = raw_body.columns = names
    above_raw = raw.loc[above_idx].copy()
    above_raw.columns = names

    # --- record the cell-level text fixes (compare raw vs cleaned)
    for col in names:
        ws_fix = ph_fix = err_fix = 0
        ex_ws = ex_ph = ""
        for r, c in zip(raw_body[col], body[col]):
            if not isinstance(r, str):
                continue
            if isnull(c):
                if r.strip().casefold() in ERROR_TOKENS:
                    err_fix += 1
                else:
                    ph_fix += r.strip() != ""
                    ex_ph = ex_ph or (f"{short(r)} -> empty" if r.strip() else "")
            elif c != r:
                ws_fix += 1
                ex_ws = ex_ws or f"{short(r)} -> {short(c)}"
        note(col, "Trimmed/normalised whitespace and hidden characters", ws_fix, ex_ws)
        note(col, "Placeholder text (N/A, null, -, ...) set to empty", ph_fix, ex_ph)
        note(col, "Excel error value (#N/A, #DIV/0!, ...) set to empty", err_fix)

    # --- drop empty rows / unnamed empty columns
    blank_rows = body.isna().all(axis=1)
    n_blank = int(blank_rows.sum())
    body, raw_body = body[~blank_rows], raw_body[~blank_rows]
    note("(sheet)", "Removed completely blank rows", n_blank)
    drop_cols = [n for n, a in zip(names, auto) if a and body[n].isna().all()]
    if drop_cols:
        note("(sheet)", "Removed empty unnamed columns", len(drop_cols))
        body, raw_body = body.drop(columns=drop_cols), raw_body.drop(columns=drop_cols)
    cols = list(body.columns)
    # above_raw keeps every column (even ones empty below the header) so no note above the table is lost
    all_cols = list(dict.fromkeys(cols + list(above_raw.columns)))
    reason_col = free_name("Reason", all_cols)
    row_col = free_name("Source Row", all_cols + [reason_col])

    if body.empty:
        warns.append(f"[{name}] has a header but no data rows.")
        rejected = pd.DataFrame()
        if len(above_raw):
            rejected = above_rows_frame(above_raw, reason_col, row_col)
        return {"name": name, "clean": body, "rejected": rejected, "log": log, "warns": warns, "fmt": {},
                "stats": dict(rows_in=0, blank=n_blank, above=len(above_raw), kept=0, invalid=0, incomplete=0, duplicate=0)}

    rej: dict = {}  # excel row -> [(tag, text)]

    def reject(idx, tag, text):
        rej.setdefault(idx, []).append((tag, text))

    new_cols, fmt = {}, {}
    lower_to_col = {c.casefold(): c for c in cols}

    # --- per-column typing
    for col in cols:
        s = body[col]
        nonnull = s.dropna()
        if nonnull.empty:
            new_cols[col] = s
            continue
        # identifier columns stay text; a heading like "Account Created Date" or "Number of Units" is not one,
        # and real Excel dates are never identifiers, so those go through date/number validation instead
        mostly_dates = sum(isinstance(v, (datetime, date, pd.Timestamp)) for v in nonnull) >= 0.5 * len(nonnull)
        id_like = (
            bool(ID_NAME_RE.search(col)) and not is_date_name(col) and not is_qty_name(col) and not mostly_dates
        ) or any(isinstance(v, str) and len(v) > 1 and v[0] == "0" and v.isdigit() for v in nonnull)
        is_email = is_email_name(col) and any(isinstance(v, str) and "@" in v for v in nonnull)
        if is_email_name(col) and not is_email:
            warns.append(f"[{name}] '{col}' looks like an email column but no value contains '@' - left unvalidated, please review.")
        # 16+ digit integers (card/account/reference numbers) lose precision as floats and could merge distinct records
        long_ints = [
            v for v in nonnull
            if (isinstance(v, int) and not isinstance(v, bool) and abs(v) >= 10**15)
            or (isinstance(v, str) and LONG_INT_RE.match(re.sub(r"[,\s$\u20ac\u00a3\u00a5]", "", v)))
        ]
        if long_ints and not is_email:
            id_like = True
            warns.append(f"[{name}] '{col}': {len(long_ints)} value(s) have 16+ digits - kept as text so no precision is lost.")

        if is_email:
            out = s.map(lambda v: v.casefold() if isinstance(v, str) else v)
            for idx, v in out.dropna().items():
                if not isinstance(v, str) or not EMAIL_RE.match(v):
                    reject(idx, "invalid", f"Invalid email in '{col}': {short(v)}")
            changed = int(((out != s) & s.notna()).sum())
            note(col, "Lower-cased email addresses", changed)
            new_cols[col] = out
            continue

        if id_like:
            def as_text(v):
                if isinstance(v, (float, np.floating)) and float(v).is_integer():
                    v = int(v)
                return str(v) if not isinstance(v, str) else v
            new_cols[col] = s.map(lambda v: None if isnull(v) else as_text(v))
            continue

        parsed = {idx: parse_number(v) for idx, v in nonnull.items()}
        num_share = sum(p is not None for p in parsed.values()) / len(parsed)

        if num_share >= 0.9:
            vals, converted, ex, any_pct = {}, 0, "", False
            for idx, v in nonnull.items():
                p = parsed[idx]
                if p is None:
                    reject(idx, "invalid", f"Non-numeric value in '{col}': {short(v)}")
                    continue
                x, pct = p
                any_pct |= pct
                vals[idx] = x
                if isinstance(v, str):
                    converted += 1
                    ex = ex or f"{short(v)} -> {x:g}"
            ser = pd.Series(vals, index=s.index, dtype="float64")
            # Int64 only when every value is a whole number that fits exactly (|x| < 2**53); else keep float
            if not any_pct and ser.dropna().map(float.is_integer).all() and ser.abs().max() < 2**53:
                ser = ser.astype("Int64")
            new_cols[col] = ser
            if any_pct:
                fmt[col] = "0.0%"
            note(col, "Converted text to numbers (currency symbols, thousands commas, %, (negatives))", converted, ex)
            if len(ser.dropna()) >= 20:
                q1, q3 = ser.dropna().astype(float).quantile([0.25, 0.75])
                iqr = q3 - q1
                if iqr > 0:
                    n_out = int(((ser.astype(float) < q1 - 3 * iqr) | (ser.astype(float) > q3 + 3 * iqr)).sum())
                    if n_out:
                        warns.append(f"[{name}] '{col}': {n_out} extreme outlier value(s) flagged but kept - please review.")
            continue

        # date detection
        is_dt = nonnull.map(lambda v: isinstance(v, (datetime, date, pd.Timestamp)))
        strs = nonnull[nonnull.map(lambda v: isinstance(v, str) and len(v) >= 6 and any(ch.isdigit() for ch in v))]
        parsed_dt = parse_dates(strs, opts.dayfirst) if len(strs) else pd.Series(dtype="datetime64[ns]")
        ok_share = (int(is_dt.sum()) + int(parsed_dt.notna().sum())) / len(nonnull)
        need = 0.5 if is_date_name(col) else 0.8
        if ok_share >= need:
            out, converted, ex = {}, 0, ""
            for idx, v in nonnull.items():
                if is_dt[idx]:
                    ts = pd.Timestamp(v)
                    if ts.tzinfo is not None:
                        ts = ts.tz_convert("UTC").tz_localize(None)
                elif idx in parsed_dt.index and not pd.isna(parsed_dt[idx]):
                    ts = parsed_dt[idx]
                    converted += 1
                    ex = ex or f"{short(v)} -> {ts:%Y-%m-%d}"
                else:
                    reject(idx, "invalid", f"Unparseable date in '{col}': {short(v)}")
                    continue
                if not 1900 <= ts.year <= 2100:
                    reject(idx, "invalid", f"Implausible date in '{col}': {short(v)}")
                    continue
                out[idx] = ts
            ser = pd.Series(out, index=s.index, dtype="datetime64[ns]")
            has_time = (ser.dropna() != ser.dropna().dt.normalize()).any()
            fmt[col] = "yyyy-mm-dd hh:mm:ss" if has_time else "yyyy-mm-dd"
            new_cols[col] = ser
            note(col, "Converted text to real Excel dates", converted, ex)
            n_tz = sum(1 for v in strs if ":" in v and TZ_RE.search(v))
            if n_tz:
                warns.append(f"[{name}] '{col}': {n_tz} timestamp(s) carried a timezone; converted to UTC and stored without one.")
            amb = [m for v in strs if (m := AMBIGUOUS_DATE_RE.match(v)) and int(m[1]) <= 12 and int(m[2]) <= 12 and m[1] != m[2]]
            if amb:
                order = "day-first" if opts.dayfirst else "month-first"
                warns.append(f"[{name}] '{col}': {len(amb)} date(s) like '{amb[0][0]}' are ambiguous; read as {order}. "
                             f"Confirm the intended format (re-run with --dayfirst if needed).")
            continue

        if 0.5 <= num_share < 0.9:
            warns.append(f"[{name}] '{col}': mixed numbers and text ({num_share:.0%} numeric) - left as text, please review.")

        # text: harmonise inconsistent casing in low-cardinality columns (e.g. 'usa' / 'USA')
        out = s
        if all(isinstance(v, str) for v in nonnull) and nonnull.nunique() <= 50 and nonnull.nunique() / len(nonnull) <= 0.5:
            groups: dict = {}
            for v in nonnull:
                groups.setdefault(v.casefold(), Counter())[v] += 1
            mapping = {}
            for cnt in groups.values():
                if len(cnt) > 1:
                    best = cnt.most_common(1)[0][0]
                    mapping.update({v: best for v in cnt if v != best})
            if mapping:
                out = s.map(lambda v: mapping.get(v, v))
                k, v0 = next(iter(mapping.items()))
                note(col, "Made inconsistent capitalisation consistent", int(s.isin(mapping).sum()), f"{short(k)} -> {short(v0)}")
        new_cols[col] = out

    clean = pd.DataFrame(new_cols, index=body.index)[cols]

    # --- completeness
    wanted = []
    for r in opts.required or []:
        if r.casefold() in lower_to_col:
            wanted.append(lower_to_col[r.casefold()])
        else:
            warns.append(f"[{name}] required column '{r}' not found - check was skipped for this sheet.")
    if wanted:
        miss = clean[wanted].isna()
        for idx in miss.index[miss.any(axis=1)]:
            reject(idx, "incomplete", "Missing required value(s): " + ", ".join(c for c in wanted if miss.at[idx, c]))
    core = [c for c in cols if clean[c].notna().mean() >= 0.5]
    if core and len(core) > 1:
        ratio = clean[core].notna().mean(axis=1)
        for idx in ratio.index[ratio < opts.min_filled]:
            filled = int(clean.loc[idx, core].notna().sum())
            reject(idx, "incomplete", f"Incomplete row: only {filled} of {len(core)} core fields filled")

    # --- duplicates (among rows that survived so far; first occurrence wins)
    def flag_dups(frame, subset, tag, label):
        if frame.empty or not subset:
            return
        gid = frame.groupby(subset, dropna=False, sort=False).ngroup()
        first = pd.Series(frame.index, index=frame.index).groupby(gid).transform("first")
        for idx in frame.index[frame.index != first]:
            reject(idx, tag, f"{label} (first seen on row {first[idx]})")

    flag_dups(clean.drop(index=list(rej)), cols, "duplicate", "Exact duplicate row")
    key_cols = [lower_to_col[k.casefold()] for k in opts.key or [] if k.casefold() in lower_to_col]
    missing_keys = [k for k in opts.key or [] if k.casefold() not in lower_to_col]
    if missing_keys:  # a partial composite key would wrongly merge distinct records, so skip entirely
        warns.append(f"[{name}] key column(s) {', '.join(repr(k) for k in missing_keys)} not found - "
                     "key-duplicate check was skipped for this sheet.")
    elif key_cols:
        live = clean.drop(index=list(rej))
        has_key = live[key_cols].notna().all(axis=1)  # a blank key is "unknown", not a match for other blanks
        if (~has_key).any():
            warns.append(f"[{name}] {int((~has_key).sum())} row(s) have a blank key ({', '.join(key_cols)}) and were not "
                         "checked for key duplicates - consider making the key required.")
        flag_dups(live[has_key], key_cols, "duplicate", f"Duplicate key ({', '.join(key_cols)})")

    # --- split kept / rejected
    bad = clean.index.isin(list(rej))
    kept = clean[~bad]
    rejected = raw_body[bad].apply(lambda c: c.map(strip_illegal))
    rejected.insert(0, reason_col, [" | ".join(t for _, t in rej[i]) for i in rejected.index])
    rejected.insert(0, row_col, rejected.index)
    if len(above_raw):  # keep rows above the header visible instead of silently dropping them
        rejected = pd.concat([above_rows_frame(above_raw, reason_col, row_col), rejected]).sort_values(row_col)

    tags = Counter(rej[i][0][0] for i in rej)
    miss_cols = {c: int(kept[c].isna().sum()) for c in cols if kept[c].isna().any()}
    if miss_cols:
        top = ", ".join(f"'{c}' ({n})" for c, n in sorted(miss_cols.items(), key=lambda x: -x[1])[:6])
        warns.append(f"[{name}] empty cells remain in kept rows: {top}. Left blank - not guessed.")

    return {
        "name": name, "clean": kept, "rejected": rejected, "log": log, "warns": warns, "fmt": fmt,
        "stats": dict(rows_in=len(body), blank=n_blank, above=len(above_raw), kept=len(kept), invalid=tags["invalid"],
                      incomplete=tags["incomplete"], duplicate=tags["duplicate"]),
    }


# ---------------------------------------------------------------- writing
def literalise(ws):
    """Store text that starts with '=' as text. openpyxl would otherwise write it as a live formula."""
    for row in ws.iter_rows():
        for c in row:
            if c.data_type == "f":
                c.data_type = "s"


def style_sheet(ws, df: pd.DataFrame, fmt: dict):
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    literalise(ws)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    ws.freeze_panes = "A2"
    if ws.max_row > 1 and ws.max_column > 0:
        ws.auto_filter.ref = ws.dimensions
    for j, col in enumerate(df.columns, start=1):
        sample = [len(str(col))] + [len(str(v)) for v in df[col].head(200) if not isnull(v)]
        ws.column_dimensions[get_column_letter(j)].width = min(max(sample) + 2, 50)
        if col in fmt:
            for row in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                row[0].number_format = fmt[col]


def write_workbook(out_path: Path, results: list, src_name: str) -> None:
    from openpyxl.styles import Font

    used = {"summary", "change log"}
    with pd.ExcelWriter(out_path, engine="openpyxl") as xw:
        sheet_names = {}
        for r in results:
            sn = unique_sheet_name(r["name"], used)
            sheet_names[r["name"]] = sn
            r["clean"].to_excel(xw, sheet_name=sn, index=False)
            style_sheet(xw.book[sn], r["clean"], r["fmt"])
        for r in results:
            if len(r["rejected"]):
                rn = unique_sheet_name(f"{r['name'][:21]} Rejected", used)
                r["rejected"].to_excel(xw, sheet_name=rn, index=False)
                style_sheet(xw.book[rn], r["rejected"], {})

        log_rows = [row for r in results for row in r["log"]]
        log_df = pd.DataFrame(log_rows, columns=["Sheet", "Column", "Action", "Count", "Example"])
        if log_df.empty:
            log_df.loc[0] = ["", "", "No automatic fixes were needed", 0, ""]
        log_df.to_excel(xw, sheet_name="Change Log", index=False)
        style_sheet(xw.book["Change Log"], log_df, {})
        xw.book["Change Log"].column_dimensions["C"].width = 70

        ws = xw.book.create_sheet("Summary", 0)
        ws.append([f"Data cleansing summary - {src_name}"])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([])
        ws.append(["Sheet", "Data rows read", "Blank rows removed", "Rows above header (moved to Rejected)", "Rows kept",
                   "Quarantined: invalid", "Quarantined: incomplete", "Quarantined: duplicate"])
        for c in ws[3]:
            c.font = Font(bold=True)
        for r in results:
            s = r["stats"]
            ws.append([sheet_names[r["name"]], s["rows_in"], s["blank"], s["above"], s["kept"],
                       s["invalid"], s["incomplete"], s["duplicate"]])
        warns = [w for r in results for w in r["warns"]]
        ws.append([])
        ws.append(["Needs your review" if warns else "No items need review"])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        for w in warns:
            ws.append([w])
        ws.append([])
        ws.append(["Quarantined rows are on the '<sheet> Rejected' sheets with the reason; nothing was silently deleted."])
        literalise(ws)
        ws.column_dimensions["A"].width = 38
        for col in "BCDEFGH":
            ws.column_dimensions[col].width = 22


# ---------------------------------------------------------------- main
def find_files(paths, recursive: bool) -> list:
    out = []
    for p in paths:
        p = Path(p)
        if p.is_dir():
            it = p.rglob("*") if recursive else p.glob("*")
            out += [f for f in it if f.is_file()]
        elif p.is_file():
            out.append(p)
        else:
            print(f"WARNING: {p} not found", file=sys.stderr)
    return sorted(
        f for f in set(out)
        if f.suffix.lower() in EXCEL_EXTS and not f.name.startswith("~$") and not f.stem.endswith("_cleaned")
    )


def output_path(path: Path, out_dir: Path | None, claimed: set) -> Path:
    """<stem>_cleaned.xlsx, or <stem>_2_cleaned.xlsx ... if another input in this run already claimed that name.
    Keeps the _cleaned.xlsx suffix so outputs are never picked up as inputs."""
    folder, n = out_dir or path.parent, 1
    while True:
        stem = path.stem if n == 1 else f"{path.stem}_{n}"
        cand = folder / f"{stem}_cleaned.xlsx"
        if str(cand.resolve()).casefold() not in claimed:
            claimed.add(str(cand.resolve()).casefold())
            return cand
        n += 1


def process(path: Path, out_dir: Path | None, opts, claimed: set) -> bool:
    print(f"\n=== {path.name}")
    try:
        sheets = read_workbook(path)
    except ImportError as e:
        print(f"ERROR: missing library for {path.suffix} files ({e}). For .xls, ask the user to re-save as .xlsx.")
        return False
    except Exception as e:  # corrupt, password-protected, not really Excel...
        print(f"ERROR: could not open file ({type(e).__name__}: {e}). It may be corrupt or password-protected.")
        return False

    results, failed = [], False
    for name, raw in sheets.items():
        try:
            res = clean_sheet(name, raw, opts)
        except HeaderReviewNeeded as e:  # needs a human decision; the other sheets are still cleaned
            print(f"ERROR: {e}")
            failed = True
            continue
        except Exception as e:  # one bad sheet must not stop the others
            print(f"ERROR: [{name}] could not be cleaned and was left out ({type(e).__name__}: {e})")
            failed = True
            continue
        if res is None:
            print(f"  [{name}] empty sheet - skipped")
        else:
            results.append(res)
    if not results:
        print("ERROR: no sheet could be cleaned (empty, or a header row needs confirming).")
        return False

    out_path = output_path(path, out_dir, claimed)
    try:
        out_path.parent.mkdir(parents=True, exist_ok=True)
        write_workbook(out_path, results, path.name)
    except Exception as e:
        print(f"ERROR: could not write the cleaned file ({type(e).__name__}: {e})")
        return False

    for r in results:
        s = r["stats"]
        print(f"  [{r['name']}] read {s['rows_in']} rows -> kept {s['kept']}, quarantined "
              f"{s['invalid']} invalid / {s['incomplete']} incomplete / {s['duplicate']} duplicate; "
              f"{s['blank']} blank rows removed")
        for row in r["log"]:
            print(f"      fix: {row['Column']}: {row['Action']} x{row['Count']}"
                  + (f"  e.g. {row['Example']}" if row["Example"] else ""))
    for r in results:
        for w in r["warns"]:
            print(f"  REVIEW: {w}")
    print(f"OUTPUT: {out_path}")
    return not failed


def main() -> int:
    ap = argparse.ArgumentParser(description="Clean Excel files; output <name>_cleaned.xlsx with a rejected-rows sheet and change log.")
    ap.add_argument("paths", nargs="*", default=["."], help="Excel files and/or folders (default: current folder)")
    ap.add_argument("--out-dir", help="Where to write cleaned files (default: next to each input)")
    ap.add_argument("--recursive", action="store_true", help="Search folders recursively")
    ap.add_argument("--key", nargs="+", metavar="COL", help="Column(s) that uniquely identify a record; later repeats are quarantined")
    ap.add_argument("--required", nargs="+", metavar="COL", help="Column(s) that must be filled; rows missing them are quarantined")
    ap.add_argument("--header-row", nargs="+", action="extend", metavar="N|SHEET=N",
                    help="Excel row number of the header: N for every sheet and/or SHEET=N for one sheet (default: auto-detect)")
    ap.add_argument("--min-filled", type=float, default=0.5, help="Quarantine rows with a smaller share of core fields filled (default 0.5)")
    ap.add_argument("--dayfirst", action="store_true", help="Read ambiguous dates as DD/MM/YYYY (default MM/DD/YYYY)")
    opts = ap.parse_args()
    for spec in opts.header_row or []:
        if not spec.rsplit("=", 1)[-1].strip().isdigit():
            ap.error(f"--header-row expects N or SHEET=N with N a row number, got '{spec}'")

    files = find_files(opts.paths, opts.recursive)
    if not files:
        print("No Excel files (.xlsx, .xlsm, .xls) found.")
        return 1
    print(f"Found {len(files)} Excel file(s).")
    claimed: set = set()
    ok = [process(f, Path(opts.out_dir) if opts.out_dir else None, opts, claimed) for f in files]
    print(f"\nDone: {sum(ok)} of {len(ok)} file(s) cleaned.")
    return 0 if all(ok) else 1


if __name__ == "__main__":
    sys.exit(main())
